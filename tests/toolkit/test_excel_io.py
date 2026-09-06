"""Pruebas unitarias para src/toolkit/excel_io.py.

Los archivos de prueba se construyen con openpyxl en un `tmp_path` real: son
`.xlsx` de verdad (hojas ocultas, celdas combinadas, filas ocultas, literales de
error), no mocks de la librería.
"""
import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from src.toolkit.excel_io import (
    coerce_excel_errors,
    excel_serial_to_datetime,
    hidden_rows_and_columns,
    inventory_workbook,
    normalize_excel_whitespace,
    profile_cell_types,
    read_sheet_expanding_merges,
)


@pytest.fixture
def messy_workbook(tmp_path):
    """Un .xlsx con los cuatro problemas de formato que este modulo resuelve."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    ws.append(["Region", "Empresa", "Produccion", "Notas"])
    ws.append(["Norte", "Codelco", 100, "borrador"])
    ws.append([None, "Escondida", 200, None])  # celda combinada con la fila anterior
    ws.append(["Sur", "Candelaria", 300, None])
    ws.append(["Sur", "Borrada", 999, None])  # fila oculta por el analista
    ws.merge_cells("A2:A3")
    ws.row_dimensions[5].hidden = True
    ws.column_dimensions["D"].hidden = True

    oculta = wb.create_sheet("Borrador")
    oculta.append(["no es dato"])
    oculta.sheet_state = "hidden"

    muy_oculta = wb.create_sheet("Parametros")
    muy_oculta.sheet_state = "veryHidden"

    path = tmp_path / "messy.xlsx"
    wb.save(path)
    return path


def test_inventory_workbook_lists_hidden_and_very_hidden_sheets(messy_workbook):
    inventory = inventory_workbook(messy_workbook)

    assert list(inventory["hoja"]) == ["Datos", "Borrador", "Parametros"]
    assert list(inventory["visible"]) == [True, False, False]
    assert inventory.loc[inventory["hoja"] == "Parametros", "estado"].item() == "veryHidden"
    assert inventory.loc[inventory["hoja"] == "Datos", "filas"].item() == 5


def test_hidden_rows_and_columns_reports_zero_based_indices(messy_workbook):
    hidden_rows, hidden_cols = hidden_rows_and_columns(messy_workbook, sheet="Datos")

    assert hidden_rows == [4]  # fila 5 de Excel
    assert hidden_cols == [3]  # columna D


def test_read_sheet_expanding_merges_fills_the_merged_range(messy_workbook):
    df = read_sheet_expanding_merges(messy_workbook, sheet="Datos")

    # "Norte" estaba solo en A2; el merge A2:A3 lo propaga a la fila siguiente.
    assert df.iloc[1, 0] == "Norte"
    assert df.iloc[2, 0] == "Norte"


def test_read_sheet_expanding_merges_keeps_hidden_rows_by_default(messy_workbook):
    df = read_sheet_expanding_merges(messy_workbook, sheet="Datos")

    assert "Borrada" in df[1].tolist()
    assert len(df) == 5


def test_read_sheet_expanding_merges_drops_each_axis_independently(messy_workbook):
    sin_filas = read_sheet_expanding_merges(messy_workbook, sheet="Datos", drop_hidden_rows=True)
    sin_columnas = read_sheet_expanding_merges(messy_workbook, sheet="Datos", drop_hidden_columns=True)

    assert (len(sin_filas), sin_filas.shape[1]) == (4, 4)
    assert "Borrada" not in sin_filas[1].tolist()
    assert (len(sin_columnas), sin_columnas.shape[1]) == (5, 3)


def test_coerce_excel_errors_replaces_literals_and_reports_them():
    df = pd.DataFrame({
        "produccion": [100, "#DIV/0!", 300, "#N/A"],
        "empresa": ["Codelco", "Escondida", "#REF!", "Candelaria"],
        "ok": [1, 2, 3, 4],
    })
    clean, report = coerce_excel_errors(df)

    assert clean["produccion"].isna().sum() == 2
    assert pd.to_numeric(clean["produccion"]).sum() == 400
    assert set(report["columna"]) == {"produccion", "empresa"}
    assert report.loc[report["columna"] == "produccion", "n_errores"].item() == 2
    assert "#REF!" in report.loc[report["columna"] == "empresa", "literales"].item()
    assert "ok" not in set(report["columna"])


def test_excel_serial_to_datetime_handles_the_1900_leap_year_bug():
    result = excel_serial_to_datetime([45292, 1, 60, 61, None])

    assert result.iloc[0] == pd.Timestamp("2024-01-01")
    assert result.iloc[1] == pd.Timestamp("1900-01-01")
    assert pd.isna(result.iloc[2])  # el 1900-02-29 que Excel inventa no existe
    assert result.iloc[3] == pd.Timestamp("1900-03-01")
    assert pd.isna(result.iloc[4])


def test_excel_serial_to_datetime_1904_system_is_offset_by_1462_days():
    mac = excel_serial_to_datetime([45292], system=1904).iloc[0]
    windows = excel_serial_to_datetime([45292]).iloc[0]

    assert (mac - windows).days == 1462


def test_normalize_excel_whitespace_removes_invisible_characters():
    df = pd.DataFrame({"empresa": ["Codelco\xa0", "​Escondida", "Los  Bronces", "   ", None]})
    clean = normalize_excel_whitespace(df)

    assert clean["empresa"].iloc[0] == "Codelco"
    assert clean["empresa"].iloc[1] == "Escondida"
    assert clean["empresa"].iloc[2] == "Los Bronces"
    assert pd.isna(clean["empresa"].iloc[3])
    assert pd.isna(clean["empresa"].iloc[4])


def test_normalize_excel_whitespace_makes_keys_join_again():
    left = pd.DataFrame({"empresa": ["Codelco\xa0", "Escondida"], "produccion": [100, 200]})
    right = pd.DataFrame({"empresa": ["Codelco", "Escondida"], "region": ["Norte", "Norte"]})

    assert left.merge(right, on="empresa").shape[0] == 1  # el espacio duro rompe el join
    assert normalize_excel_whitespace(left).merge(right, on="empresa").shape[0] == 2


def test_profile_cell_types_flags_numbers_stored_as_text():
    df = pd.DataFrame({
        "mixta": [1.0, 2.0, "3,5", np.nan],
        "limpia": [1.0, 2.0, 3.0, 4.0],
    })
    profile = profile_cell_types(df).set_index("columna")

    assert bool(profile.loc["mixta", "mixta"]) is True
    assert profile.loc["mixta", "n_texto"] == 1
    assert profile.loc["mixta", "n_numero"] == 2
    assert profile.loc["mixta", "n_nulos"] == 1
    assert profile.loc["mixta", "tipo_dominante"] == "numero"
    assert bool(profile.loc["limpia", "mixta"]) is False
